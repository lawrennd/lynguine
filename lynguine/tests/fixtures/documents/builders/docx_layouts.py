"""Synthetic DOCX builders for dialogue strategies (CIP-0013)."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument


def make_chatbot_transcript(path: str | Path) -> Path:
    """Minimal chatbot export shape (Participant / Transcript / turns)."""
    path = Path(path)
    doc = DocxDocument()
    doc.add_paragraph("Participant 001")
    doc.add_paragraph("Conversation Summary")
    doc.add_paragraph("A short summary of the chat.")
    doc.add_paragraph("Conversation Transcript")
    doc.add_paragraph("What concerns you about AI in local government?")
    doc.add_paragraph(
        "Participant 001: I worry about accountability when decisions are automated "
        "without a clear human owner."
    )
    doc.add_paragraph("What benefits do you see?")
    doc.add_paragraph(
        "Participant 001: Faster responses for routine enquiries could free staff time."
    )
    doc.add_paragraph("Participant 002")
    doc.add_paragraph("Conversation Transcript")
    doc.add_paragraph("What concerns you?")
    doc.add_paragraph(
        "Participant 002: Privacy of personal data used to train systems is important."
    )
    doc.save(path)
    return path


def make_workshop_transcript(path: str | Path) -> Path:
    """Workshop paragraphs with speaker labels for chunking."""
    path = Path(path)
    doc = DocxDocument()
    doc.add_paragraph("Facilitator: Welcome everyone to the workshop on AI.")
    body = (
        "Speaker 1: I think councils should publish when they use automated tools. "
        "Speaker 2: Yes and there should be a way to appeal a decision. "
    )
    doc.add_paragraph(body * 8)
    doc.add_paragraph(
        "Facilitator: Thank you. Let's hear about benefits as well. " + body * 4
    )
    doc.save(path)
    return path
