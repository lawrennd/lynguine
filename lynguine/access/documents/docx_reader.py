"""DOCX reader via python-docx (optional) into the document model (CIP-0013)."""

from __future__ import annotations

from typing import Any

from lynguine.access.paths import confine_configured_path, resolve_under_roots

from .model import Block, Document, Page, TextLayer
from .pdf import classify_text_layer

_DOCX_INSTALL = (
    "python-docx is required for DOCX reading. "
    "Install with: pip install 'lynguine[documents]' "
    "or: poetry install --extras documents"
)


class PythonDocxMissingError(ImportError):
    """Raised when DOCX reading is requested but python-docx is not installed."""


def _import_docx():
    try:
        from docx import Document as DocxDocument  # type: ignore
    except ImportError as exc:
        raise PythonDocxMissingError(_DOCX_INSTALL) from exc
    return DocxDocument


def read_docx(
    path: str,
    *,
    allowed_roots: list[str] | None = None,
    details: dict[str, Any] | None = None,
) -> Document:
    """Open a DOCX under path confinement and return a :class:`Document`.

    Each non-empty paragraph becomes a text block on a single logical page
    (Word has no fixed page geometry here).
    """
    DocxDocument = _import_docx()

    if details is not None:
        resolved = confine_configured_path(path, details)
    elif allowed_roots is not None:
        resolved = resolve_under_roots(path, allowed_roots)
    else:
        resolved = path

    docx = DocxDocument(resolved)
    blocks: list[Block] = []
    texts: list[str] = []
    for i, para in enumerate(docx.paragraphs):
        text = (para.text or "").replace("\r\n", "\n").replace("\r", "\n").replace(
            "\xa0", " "
        )
        if not text.strip():
            continue
        texts.append(text)
        blocks.append(
            Block(
                page_index=0,
                block_index=i,
                x0=0.0,
                y0=float(i),
                x1=1.0,
                y1=float(i) + 1.0,
                text=text.strip(),
                block_type="text",
                word_count=len(text.split()),
            )
        )
    full = "\n".join(texts)
    total_words = len(full.split()) if full.strip() else 0
    text_layer, words_per_page = classify_text_layer(total_words, page_count=1)
    page = Page(
        index=0,
        width=1.0,
        height=float(max(len(blocks), 1)),
        text=full,
        word_count=total_words,
        blocks=tuple(blocks),
    )
    return Document(
        path=resolved,
        reader="docx",
        page_count=1,
        total_words=total_words,
        words_per_page=words_per_page,
        text_layer=text_layer,
        pages=(page,),
        metadata={"paragraph_count": len(blocks)},
    )
